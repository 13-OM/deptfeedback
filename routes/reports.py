import re
from io import BytesIO
from datetime import datetime
import pandas as pd
from flask import Blueprint,abort,request,send_file
from flask_login import current_user
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
from openpyxl.utils import get_column_letter
from decorators import role_required
from models import FeedbackForm,FeedbackQuestion,Subject
from services import average_rating,completion_for_forms,responses_for_forms
bp=Blueprint("reports",__name__,url_prefix="/reports")
def _safe_name(v):return re.sub(r"[^A-Za-z0-9_-]+","_",v or "DeptFeedback").strip("_")[:55] or "DeptFeedback"
def _excel_safe(v):return "'"+v if isinstance(v,str) and (v.startswith(("=","+","-","@")) or v.startswith((chr(9),chr(13)))) else v
def _build_workbook(forms,responses):
    questions=FeedbackQuestion.all({},[("question_order",1)]); rating_questions=[q for q in questions if not q.is_comment]; response_records=[];by_form={};by_faculty={};qvals={q.id:[] for q in rating_questions}
    for r in responses:
        f=r.feedback_form;s=f.subject;faculty=s.faculty.name if s.faculty else "Unassigned";am={a.question_id:a for a in r.answers};row={"Feedback ID":r.anonymous_reference,"Subject":s.subject_name,"Subject Code":s.subject_code,"Semester":s.semester,"Faculty":faculty,"Academic Year":f.academic_year};ratings=[]
        for i,q in enumerate(rating_questions,1):
            a=am.get(q.id);rating=a.rating if a else None;row[f"Q{i} — {q.question_text}"]=rating
            if rating is not None:ratings.append(int(rating));qvals[q.id].append(int(rating))
        comments=[(a.comment or "").strip() for a in r.answers if a.question and a.question.is_comment and (a.comment or "").strip()];row["Overall Rating"]=round(sum(ratings)/len(ratings),2) if ratings else None;row["Comment"]=" | ".join(comments);row["Submitted At"]=r.submitted_at.strftime("%Y-%m-%d %H:%M") if r.submitted_at else "";response_records.append(row);by_form.setdefault(f.id,[]).append(r);by_faculty.setdefault(faculty,[]).append((s,r))
    cols=["Feedback ID","Subject","Subject Code","Semester","Faculty","Academic Year"]+[f"Q{i} — {q.question_text}" for i,q in enumerate(rating_questions,1)]+["Overall Rating","Comment","Submitted At"]
    responses_df=pd.DataFrame([{k:_excel_safe(v) for k,v in x.items()} for x in response_records],columns=cols)
    subject_records=[]
    for f in forms:
        rs=by_form.get(f.id,[]);s=f.subject;subject_records.append({"Subject":s.subject_name,"Subject Code":s.subject_code,"Semester":s.semester,"Academic Year":f.academic_year,"Faculty":s.faculty.name if s.faculty else "Unassigned","Responses":len(rs),"Average Rating":average_rating(rs)})
    subject_df=pd.DataFrame(subject_records,columns=["Subject","Subject Code","Semester","Academic Year","Faculty","Responses","Average Rating"])
    faculty_records=[]
    for name,pairs in sorted(by_faculty.items()):
        rs=[r for _,r in pairs];faculty_records.append({"Faculty":name,"Subjects":", ".join(sorted({s.subject_name for s,_ in pairs})),"Responses":len(rs),"Average Rating":average_rating(rs)})
    faculty_df=pd.DataFrame(faculty_records,columns=["Faculty","Subjects","Responses","Average Rating"])
    qdf=pd.DataFrame([{"Question":q.question_text,"Category":q.category,"Responses":len(qvals[q.id]),"Average Rating":round(sum(qvals[q.id])/len(qvals[q.id]),2) if qvals[q.id] else 0} for q in rating_questions])
    comp=completion_for_forms(forms);vals=[a.rating for r in responses for a in r.answers if a.rating is not None];stats=pd.DataFrame([{"Metric":"Eligible students across selected forms","Value":comp["eligible"]},{"Metric":"Anonymous feedback responses","Value":len(responses)},{"Metric":"Feedback completion rate","Value":f"{comp['rate']:.1f}%"},{"Metric":"Overall average rating","Value":f"{sum(vals)/len(vals):.2f} / 5" if vals else "—"},{"Metric":"Report generated","Value":datetime.now().strftime("%Y-%m-%d %H:%M")},{"Metric":"Privacy note","Value":"No student identity is included in response reports."}])
    buf=BytesIO()
    with pd.ExcelWriter(buf,engine="openpyxl") as w:
        responses_df.to_excel(w,sheet_name="Feedback Responses",index=False);subject_df.to_excel(w,sheet_name="Subject Summary",index=False);faculty_df.to_excel(w,sheet_name="Faculty Summary",index=False);qdf.to_excel(w,sheet_name="Question Analysis",index=False);stats.to_excel(w,sheet_name="Feedback Statistics",index=False)
    buf.seek(0);wb=load_workbook(buf);fill=PatternFill("solid",fgColor="172554");font=Font(name="Aptos",bold=True,color="FFFFFF",size=10);body=Font(name="Aptos",size=10,color="243047");line=Side(style="thin",color="E2E8F0")
    for sh in wb.worksheets:
        sh.freeze_panes="A2";sh.auto_filter.ref=sh.dimensions;sh.sheet_view.showGridLines=False;sh.row_dimensions[1].height=30
        for c in sh[1]:c.fill=fill;c.font=font;c.alignment=Alignment(vertical="center",wrap_text=True)
        for row in sh.iter_rows(min_row=2):
            for c in row:c.font=body;c.alignment=Alignment(vertical="top",wrap_text=True);c.border=Border(bottom=line)
        for cells in sh.columns:
            letter=get_column_letter(cells[0].column);header=str(cells[0].value or "");vals=[str(c.value or "") for c in cells[1:30]];mx=max([len(header)]+[min(len(v),48) for v in vals]);sh.column_dimensions[letter].width=42 if header in {"Comment","Question"} else 34 if header.startswith("Q") and " — " in header else 34 if header=="Metric" else 54 if header=="Value" else max(12,min(mx+2,28))
    out=BytesIO();wb.save(out);out.seek(0);return out
@bp.get("/export.xlsx")
@role_required("faculty","hod")
def export_excel():
    if current_user.role=="faculty":
        faculty=current_user.faculty;forms=[f for f in FeedbackForm.all() if f.subject and f.subject.faculty_id==faculty.id];subject_filter=request.args.get("subject","").strip()
        if subject_filter.isdigit():
            subject=Subject.get(int(subject_filter));
            if not subject or subject.faculty_id!=faculty.id:abort(403)
            forms=[f for f in forms if f.subject_id==subject.id]
        filename=f"Faculty_Feedback_{_safe_name(faculty.name)}.xlsx"
    else:
        from routes.hod import filtered_forms
        forms=filtered_forms(request.args);filename=f"Department_Feedback_Sem{_safe_name(request.args.get('semester','All') or 'All')}_{_safe_name(request.args.get('academic_year','All') or 'All')}.xlsx"
    return send_file(_build_workbook(forms,responses_for_forms(forms)),as_attachment=True,download_name=filename,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",max_age=0)
