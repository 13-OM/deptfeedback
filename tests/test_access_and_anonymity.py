import unittest
from io import BytesIO

from openpyxl import load_workbook

from app import create_app
from extensions import db
from models import FeedbackAnswer, FeedbackForm, FeedbackQuestion, FeedbackResponse, Student, Subject, SubmissionTracking


class DeptFeedbackAccessTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = create_app({
            "TESTING": True,
            "WTF_CSRF_ENABLED": False,
            "SECRET_KEY": "test-secret",
            "SQLALCHEMY_DATABASE_URI": "sqlite://",
            "SEED_DEMO": True,
        })

    def setUp(self):
        self.client = self.app.test_client()

    def login(self, role, username, password):
        return self.client.post("/login", data={
            "role": role, "username": username, "password": password,
        }, follow_redirects=True)

    def logout(self):
        self.client.post("/logout")

    def test_role_routes_and_faculty_scope_are_server_enforced(self):
        with self.app.app_context():
            own = FeedbackForm.query.join(Subject).filter(Subject.subject_name == "Python Programming").first()
            other = Subject.query.filter_by(subject_name="Web Application Development").first()
            own_id, other_id = own.subject_id, other.id

        self.login("student", "250183107002", "student123")
        self.assertEqual(self.client.get("/hod/analytics").status_code, 403)
        self.logout()

        self.login("faculty", "faculty1", "faculty123")
        self.assertEqual(self.client.get(f"/faculty/subject/{own_id}").status_code, 200)
        self.assertEqual(self.client.get(f"/faculty/subject/{other_id}").status_code, 403)
        self.assertEqual(self.client.get(f"/reports/export.xlsx?subject={other_id}").status_code, 403)

    def test_submission_is_anonymous_and_duplicate_is_blocked(self):
        with self.app.app_context():
            form = FeedbackForm.query.join(Subject).filter(Subject.subject_name == "Python Programming").first()
            form_id = form.id
            question_ids = [q.id for q in FeedbackQuestion.query.filter_by(status="Active", is_comment=False).all()]
            before = FeedbackResponse.query.filter_by(feedback_form_id=form_id).count()

        self.login("student", "250183107002", "student123")
        payload = {f"rating_{question_id}": "4" for question_id in question_ids}
        payload["comment"] = "More guided examples would help."
        first = self.client.post(f"/student/feedback/{form_id}", data=payload, follow_redirects=True)
        self.assertEqual(first.status_code, 200)
        self.assertIn(b"Feedback submitted successfully", first.data)

        with self.app.app_context():
            student_id = Student.query.filter_by(enrollment_no="250183107002").first().id
            after = FeedbackResponse.query.filter_by(feedback_form_id=form_id).count()
            self.assertEqual(after, before + 1)
            self.assertEqual(SubmissionTracking.query.filter_by(student_id=student_id, feedback_form_id=form_id).count(), 1)
            self.assertNotIn("student_id", FeedbackResponse.__table__.columns.keys())
            self.assertNotIn("enrollment_no", FeedbackResponse.__table__.columns.keys())
            self.assertNotIn("user_id", FeedbackResponse.__table__.columns.keys())
            response = FeedbackResponse.query.filter_by(feedback_form_id=form_id).order_by(FeedbackResponse.id.desc()).first()
            self.assertTrue(any(answer.comment for answer in response.answers))
            self.assertFalse(hasattr(response, "student_id"))

        second = self.client.post(f"/student/feedback/{form_id}", data=payload, follow_redirects=True)
        self.assertEqual(second.status_code, 200)
        self.assertIn(b"already submitted", second.data.lower())
        with self.app.app_context():
            self.assertEqual(FeedbackResponse.query.filter_by(feedback_form_id=form_id).count(), before + 1)

    def test_excel_report_has_expected_sheets_and_no_student_columns(self):
        self.login("hod", "hod", "hod123")
        response = self.client.get("/reports/export.xlsx?semester=5&academic_year=2026-27")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        self.assertEqual(workbook.sheetnames, [
            "Feedback Responses", "Subject Summary", "Faculty Summary",
            "Question Analysis", "Feedback Statistics",
        ])
        headers = [cell.value for cell in next(workbook["Feedback Responses"].iter_rows(min_row=1, max_row=1))]
        self.assertIn("Feedback ID", headers)
        self.assertNotIn("Student ID", headers)
        self.assertNotIn("Enrollment Number", headers)
        self.assertFalse(any("enrollment" in str(value).lower() for value in headers))


if __name__ == "__main__":
    unittest.main()
